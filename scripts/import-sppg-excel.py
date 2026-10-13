"""Import only workbook values; spreadsheet cells are data, never instructions."""
import argparse, collections, datetime, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook
parser=argparse.ArgumentParser()
parser.add_argument('daily');parser.add_argument('pm');parser.add_argument('--as-of',required=True)
args=parser.parse_args();cutoff=datetime.date.fromisoformat(args.as_of)
root=Path(__file__).resolve().parents[1]
norm=lambda text:re.sub(r'[^a-z0-9]','',str(text).lower())
def val(row,index):return row[index] if len(row)>index else None
def text(value):return str(value).strip() if value is not None else ''
def number(value):
 if isinstance(value,(int,float)) and value>=0 and float(value).is_integer():return int(value)
 return None
def write(name,value): (root/'data'/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')
seed=json.loads((root/'data/sppg.json').read_text());byname={norm(r['nama']):r for r in seed}
statuses={'OPERASIONAL':'operasional','AKAN OPERASIONAL':'akan','BERHENTI OPS SEMENTARA':'berhenti','BERHENTI OPS VA':'berhenti','SUSPEND':'suspend'}
book=load_workbook(args.daily,read_only=True,data_only=True)
history=[];latest={};positive={};unmatched=[];mismatches=[]
for sheet in book:
 match=re.fullmatch(r'(\d+) OKT',sheet.title.strip())
 if not match:continue
 date=datetime.date(cutoff.year,10,int(match[1]))
 if date>cutoff:continue
 for index,row in enumerate(sheet.iter_rows(values_only=True),1):
  name=text(val(row,1))
  if not name.upper().startswith('SPPG '):continue
  target=byname.get(norm(name))
  if target is None:unmatched.append({'sheet':sheet.title,'row':index,'nama':name});continue
  status=text(val(row,3));total=number(val(row,6))
  parts={text(c.value).replace('\n',' '):number(val(row,c.column-1)) for c in list(sheet.iter_rows(min_row=2,max_row=2))[0][7:34] if c.value is not None}
  active=bool(text(val(row,4))) or any((v or 0)>0 for v in parts.values()) or (total or 0)>0
  record={'sppgId':target['id'],'tanggal':date.isoformat(),'statusLabel':status,'namaKa':text(val(row,2)),'menu':text(val(row,4)),'jumlahKpm':number(val(row,5)),'totalPm':total if active else None,'rincian':parts,'dokumentasi':text(val(row,34)),'keterangan':text(val(row,35)),'sourceSheet':sheet.title,'sourceRow':index}
  history.append(record)
  if target['id'] not in latest or record['tanggal']>=latest[target['id']]['tanggal']:latest[target['id']]=record
  if (total or 0)>0 and (target['id'] not in positive or record['tanggal']>=positive[target['id']]['tanggal']):positive[target['id']]=record
  known=[v for v in parts.values() if v is not None]
  if (total or 0)>0 and sum(known)!=total:mismatches.append({'sppgId':target['id'],'tanggal':date.isoformat(),'total':total,'sumParts':sum(known)})
if unmatched:raise SystemExit('Unmatched SPPG rows; import aborted')
updated=0
for target in seed:
 latestrow=latest.get(target['id'])
 if latestrow:
  label=latestrow['statusLabel']
  if label not in statuses:raise SystemExit('Unrecognized status; import aborted')
  target['status']=statuses[label];target.setdefault('detail',{})['statusLabel']=label
  target['detail']['statusAsOf']=latestrow['tanggal']
  ka=latestrow['namaKa']
  if ka:
   previous=target['detail'].get('namaKa','')
   if previous and norm(previous)!=norm(ka):target['detail']['kaContactNeedsVerification']=True
   target['detail']['namaKa']=ka
 row=positive.get(target['id'])
 if not row:continue
 target['porsi']=row['totalPm'];d=target.setdefault('detail',{});parts=row['rincian']
 d['pmAsOf']='Laporan harian '+datetime.date.fromisoformat(row['tanggal']).strftime('%d')+' Oktober '+str(cutoff.year)
 d['jenjang']={k:v for k,v in parts.items() if v is not None};d['dailyReportDate']=row['tanggal'];d['jumlahKpm']=row['jumlahKpm']
 vals=lambda k:parts.get(k) or 0
 balita=vals('Balita Non PAUD');bumil=vals('BUMIL');busui=vals('BUSUI');guru=vals('Guru & Tendik')
 students=sum(v or 0 for k,v in parts.items() if k not in ['Balita Non PAUD','BUMIL','BUSUI','Guru & Tendik'])
 pm=target.setdefault('pm',{})
 for k in ['pmSatdikKelompok','pm3bKelompok']:pm.pop(k,None)
 pm.update(pmSiswa=students,pmGuruTendik=guru,pmSatdikTotal=students+guru,pmBalita=balita,pmBumil=bumil,pmBusui=busui,pm3bTotal=balita+bumil+busui)
 d['pmReconciled']=sum(v or 0 for v in parts.values())==row['totalPm']
 updated+=1
pmbook=load_workbook(args.pm,read_only=True,data_only=True);groups=[];keys=collections.Counter();missing=[]
for sheet in pmbook:
 for rownum,row in enumerate(sheet.iter_rows(values_only=True),1):
  name=text(val(row,11));label=text(val(row,0))
  if not name or name=='Nama Kelompok' or label.lower()=='contoh':continue
  kec=text(val(row,5)) or sheet.title.replace('KECAMATAN','').strip()
  if kec=='....':continue
  npsn=text(val(row,9));npsn=npsn.removesuffix('.0')
  sums=[number(val(row,i)) for i in [17,18,31,32,33,34,35]]
  total=number(val(row,36));computed=sum(v or 0 for v in sums)
  key=norm(kec)+'|'+norm(name)+'|'+npsn
  keys[key]+=1
  r={'id':f'pm-{norm(sheet.title)}-{rownum}','kecamatan':kec.title(),'desa':text(val(row,6)),'alamat':text(val(row,8)),'kodePos':text(val(row,7)).removesuffix('.0'),'lat':val(row,1) if isinstance(val(row,1),(int,float)) and -90<=val(row,1)<=90 else None,'lng':val(row,2) if isinstance(val(row,2),(int,float)) and -180<=val(row,2)<=180 else None,'nama':name,'jenis':text(val(row,10)),'npsn':npsn,'statusPelayanan':text(val(row,13)) or 'Belum diisi','optOut':number(val(row,12)),'pmSiswaLaki':sums[0],'pmSiswaPerempuan':sums[1],'pmPendidik':sums[2],'pmTendik':sums[3],'pmBumil':sums[4],'pmBusui':sums[5],'pmBalita':sums[6],'totalPm':total,'sumParts':computed,'reconciled':total==computed if total is not None else None,'sourceSheet':sheet.title,'sourceRow':rownum,'groupKey':key}
  groups.append(r)
for r in groups:r['possibleDuplicate']=keys[r.pop('groupKey')]>1
metadata={'importAsOf':cutoff.isoformat(),'dailyWorkbook':Path(args.daily).name,'pmWorkbook':Path(args.pm).name,'dailySha256':hashlib.sha256(Path(args.daily).read_bytes()).hexdigest(),'pmSha256':hashlib.sha256(Path(args.pm).read_bytes()).hexdigest(),'sppgCount':len(seed),'updatedPmCount':updated,'retainedPriorPmCount':len(seed)-updated,'latestStatuses':dict(collections.Counter(r['status'] for r in seed)),'dailyRecords':len(history),'dailyTotalMismatches':len(mismatches),'pmGroupRows':len(groups),'possibleDuplicateRows':sum(r['possibleDuplicate'] for r in groups),'pmTotalMismatches':sum(r['reconciled'] is False for r in groups),'notes':['Dates after cutoff excluded. Blank/default zero daily totals do not overwrite last positive report.','PM workbook has no SPPG foreign key. Do not join groups to kitchens by kecamatan.','Example/template rows excluded; valid named rows with blank row numbers retained.','NIK/PIC phone fields omitted from PM analytics because they are unnecessary for aggregate reporting.','Totals kept as reported; inconsistent subtotals and possible duplicate groups flagged, not silently reconciled.']}
write('sppg.json',seed);write('laporan-oktober.json',history);write('pm-kabupaten.json',groups);write('sppg-import-summary.json',metadata)
print(json.dumps(metadata,ensure_ascii=False,indent=2))
